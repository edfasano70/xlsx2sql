from __future__ import annotations

import re
import sys
import xml.etree.ElementTree as ET
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook
from PyQt6.QtCore import Qt
from PyQt6.QtGui import QAction, QIcon, QPixmap
from PyQt6.QtWidgets import (
    QAbstractItemView,
    QApplication,
    QDialog,
    QFileDialog,
    QHBoxLayout,
    QLabel,
    QListWidget,
    QMainWindow,
    QMessageBox,
    QPlainTextEdit,
    QPushButton,
    QVBoxLayout,
    QWidget,
)

PROGRAM_NAME = "xlsx2sql"
PROGRAM_VERSION = "0.2.1"
PROGRAM_DESCRIPTION = "Convierte archivos Excel (.xlsx) en sentencias SQL INSERT"

RESOURCES_DIR = Path(__file__).resolve().parent / "resources"
ICONS_DIR = RESOURCES_DIR / "icons"
LOGO_FILE = RESOURCES_DIR / "logo.png"

CONFIG_DIR = Path.home() / ".config" / PROGRAM_NAME
CONFIG_FILE = CONFIG_DIR / "config.xml"

# Identificadores SQL seguros: letras, números y guiones bajos (empezando por letra o _)
_IDENT_RE = re.compile(r"^[A-Za-z_][A-Za-z0-9_]*$")

DARK_STYLE = """
    QMainWindow, QWidget, QDialog { background-color: #1e1e1e; color: #e0e0e0; }
    QMenuBar { background-color: #2d2d2d; color: #e0e0e0; }
    QMenuBar::item { padding: 4px 8px; }
    QMenuBar::item:selected { background-color: #3d3d3d; }
    QMenu { background-color: #2d2d2d; color: #e0e0e0; }
    QMenu::item:selected { background-color: #FF9800; color: #1e1e1e; }
    QLabel { color: #e0e0e0; }
    QPushButton { background-color: #3d3d3d; border: 1px solid #555;
        border-radius: 3px; padding: 4px 8px; color: #e0e0e0; }
    QPushButton:hover { background-color: #4d4d4d; }
    QPlainTextEdit { background-color: #252525; color: #e0e0e0;
        border: 1px solid #555; font-family: monospace; }
"""

INSTRUCTIONS_TEXT = (
    "1. Pulsa «Abrir Excel...» (Ctrl+O) y selecciona un archivo .xlsx.\n"
    "2. Elige en la ventana emergente la tabla (hoja) que quieras convertir.\n"
    "3. Revisa el SQL generado; puedes editarlo.\n"
    "4. Pulsa «Guardar SQL...» (Ctrl+S) para guardarlo en un archivo .sql.\n\n"
    "Marca «Opciones > Incluir CREATE TABLE» para generar también la sentencia "
    "CREATE TABLE.\n"
    "Cada hoja se convierte en una tabla con el mismo nombre. "
    "Las filas completamente vacías se omiten."
)


def quote_identifier(name: str) -> str:
    """Devuelve el identificador tal cual si es seguro; si no, lo entrecomilla.

    Usa comillas dobles estándar SQL y duplica las comillas internas.
    """
    name = str(name)
    if _IDENT_RE.match(name):
        return name
    return '"' + name.replace('"', '""') + '"'


def sql_value(value, data_type: str) -> str:
    """Convierte el valor de una celda a un literal SQL según su tipo real."""
    if value is None:
        return "NULL"

    # Strings (incluye strings "inline" de Excel): escapar comillas simples
    if data_type == "s":
        return "'" + str(value).replace("'", "''") + "'"

    # Booleanos
    if data_type == "b":
        return "1" if value else "0"

    # Números
    if data_type == "n":
        # Evitar notación de punto flotante para enteros
        if isinstance(value, float) and value.is_integer():
            value = int(value)
        return str(value)

    # Fechas / datetimes
    if data_type == "d" or isinstance(value, (datetime, date)):
        return "'" + value.isoformat() + "'"

    # Fallback: tratar como texto escapando comillas simples
    return "'" + str(value).replace("'", "''") + "'"


def list_sheet_names(excel_file_path: str) -> list[str]:
    """Devuelve los nombres de las hojas (tablas) del archivo Excel."""
    workbook = load_workbook(excel_file_path, read_only=True)
    return workbook.sheetnames


def _column_type(cells: list[tuple]) -> str:
    """Infiera un tipo SQL simple a partir de las celdas de una columna."""
    has_data = False
    all_integer = True
    all_numeric = True
    all_boolean = True
    all_date = True

    for value, data_type in cells:
        if value is None:
            continue
        has_data = True
        if data_type == "b":
            all_integer = False
            all_numeric = False
            all_date = False
        elif data_type == "n":
            all_boolean = False
            all_date = False
            if isinstance(value, float) and not value.is_integer():
                all_integer = False
        elif data_type == "d" or isinstance(value, (datetime, date)):
            all_integer = False
            all_numeric = False
            all_boolean = False
        else:
            all_integer = False
            all_numeric = False
            all_boolean = False
            all_date = False

    if not has_data:
        return "TEXT"
    if all_boolean:
        return "BOOLEAN"
    if all_integer:
        return "INTEGER"
    if all_numeric:
        return "REAL"
    if all_date:
        return "DATE"
    return "TEXT"


def generate_insert_statements(
    excel_file_path: str,
    sheet_names: list[str] | None = None,
    include_create: bool = False,
) -> list[str]:
    """Lee un Excel y devuelve la lista de sentencias INSERT.

    Si `sheet_names` se indica, solo se convierten esas hojas.
    Si `include_create` es True, se antepone un CREATE TABLE IF NOT EXISTS por hoja.

    Usa openpyxl directamente para respetar el tipo real de cada celda
    (pandas convertiría códigos como phone/token a números y perdería
    los ceros a la izquierda).
    """
    workbook = load_workbook(excel_file_path, data_only=True)
    statements: list[str] = []

    for worksheet in workbook.worksheets:
        if sheet_names is not None and worksheet.title not in sheet_names:
            continue
        rows = list(worksheet.iter_rows())
        if not rows:
            continue

        # La primera fila son los nombres de columna
        headers = rows[0]

        # Solo se usan las columnas cuyo encabezado no está vacío
        valid_indexes = [
            i
            for i, cell in enumerate(headers)
            if cell.value is not None and str(cell.value).strip() != ""
        ]
        if not valid_indexes:
            continue

        table = quote_identifier(worksheet.title)
        columns = ", ".join(quote_identifier(headers[i].value) for i in valid_indexes)

        if include_create:
            column_defs = []
            for i in valid_indexes:
                cells = [(row[i].value, row[i].data_type) for row in rows[1:]]
                column_defs.append(f"{quote_identifier(headers[i].value)} {_column_type(cells)}")
            statements.append(f"CREATE TABLE IF NOT EXISTS {table} ({', '.join(column_defs)});")

        for row in rows[1:]:
            values = [sql_value(row[i].value, row[i].data_type) for i in valid_indexes]

            # Descartar filas completamente vacías para no generar INSERTs basura
            if all(value == "NULL" for value in values):
                continue

            statements.append(f"INSERT INTO {table} ({columns}) VALUES ({', '.join(values)});")

    return statements


def _icon(name: str) -> QIcon:
    path = ICONS_DIR / f"{name}.svg"
    return QIcon(str(path)) if path.is_file() else QIcon()


def _logo_icon() -> QIcon:
    return QIcon(str(LOGO_FILE)) if LOGO_FILE.is_file() else QIcon()


class AboutDialog(QDialog):
    def __init__(self, parent: QWidget | None = None, dark_mode: bool = False) -> None:
        super().__init__(parent)
        self.setWindowTitle("Acerca de")
        self.setWindowIcon(_logo_icon())
        self.setWindowFlags(self.windowFlags() & ~Qt.WindowType.WindowContextHelpButtonHint)

        layout = QVBoxLayout(self)

        logo_label = QLabel()
        logo_label.setAlignment(Qt.AlignmentFlag.AlignCenter)
        if LOGO_FILE.is_file():
            pixmap = QPixmap(str(LOGO_FILE))
            if not pixmap.isNull():
                pixmap = pixmap.scaled(
                    120,
                    120,
                    Qt.AspectRatioMode.KeepAspectRatio,
                    Qt.TransformationMode.SmoothTransformation,
                )
                logo_label.setPixmap(pixmap)
        layout.addWidget(logo_label)

        title_label = QLabel(f"<b>{PROGRAM_NAME}</b>")
        title_label.setAlignment(Qt.AlignmentFlag.AlignCenter)
        layout.addWidget(title_label)

        version_label = QLabel(f"Versión {PROGRAM_VERSION}")
        version_label.setAlignment(Qt.AlignmentFlag.AlignCenter)
        layout.addWidget(version_label)

        description_label = QLabel(PROGRAM_DESCRIPTION)
        description_label.setWordWrap(True)
        layout.addWidget(description_label)

        author_label = QLabel("Programado por @EdFasano70")
        author_label.setAlignment(Qt.AlignmentFlag.AlignCenter)
        layout.addWidget(author_label)

        close_button = QPushButton("Cerrar")
        close_button.clicked.connect(self.accept)
        layout.addWidget(close_button)

        if dark_mode:
            self.setStyleSheet(DARK_STYLE)


class HelpDialog(QDialog):
    def __init__(self, parent: QWidget | None = None, dark_mode: bool = False) -> None:
        super().__init__(parent)
        self.setWindowTitle("Instrucciones")
        self.setWindowIcon(_icon("help"))
        self.setWindowFlags(self.windowFlags() & ~Qt.WindowType.WindowContextHelpButtonHint)

        layout = QVBoxLayout(self)

        text_label = QLabel(INSTRUCTIONS_TEXT)
        text_label.setWordWrap(True)
        layout.addWidget(text_label)

        close_button = QPushButton("Cerrar")
        close_button.clicked.connect(self.accept)
        layout.addWidget(close_button)

        if dark_mode:
            self.setStyleSheet(DARK_STYLE)


class SheetSelectionDialog(QDialog):
    def __init__(
        self,
        sheet_names: list[str],
        parent: QWidget | None = None,
        dark_mode: bool = False,
    ) -> None:
        super().__init__(parent)
        self.selected_sheet: str | None = None
        self.setWindowTitle("Seleccionar tabla")
        self.setWindowIcon(_icon("build"))
        self.setWindowFlags(self.windowFlags() & ~Qt.WindowType.WindowContextHelpButtonHint)

        layout = QVBoxLayout(self)

        label = QLabel("Selecciona la tabla a convertir:")
        layout.addWidget(label)

        self.sheet_list = QListWidget()
        self.sheet_list.addItems(sheet_names)
        self.sheet_list.setSelectionMode(QAbstractItemView.SelectionMode.SingleSelection)
        if sheet_names:
            self.sheet_list.setCurrentRow(0)
        self.sheet_list.itemDoubleClicked.connect(lambda _item: self._accept())
        layout.addWidget(self.sheet_list)

        buttons = QHBoxLayout()
        cancel_button = QPushButton("Cancelar")
        cancel_button.clicked.connect(self.reject)
        buttons.addWidget(cancel_button)

        ok_button = QPushButton("Convertir")
        ok_button.setDefault(True)
        ok_button.clicked.connect(self._accept)
        buttons.addWidget(ok_button)

        layout.addLayout(buttons)

        if dark_mode:
            self.setStyleSheet(DARK_STYLE)

    def _accept(self) -> None:
        item = self.sheet_list.currentItem()
        if item is not None:
            self.selected_sheet = item.text()
        self.accept()


class MainWindow(QMainWindow):
    def __init__(self) -> None:
        super().__init__()
        self.dark_mode = False
        self.include_create = False
        self.excel_path: Path | None = None

        self.setWindowTitle(PROGRAM_NAME)
        self.setWindowIcon(_logo_icon())
        self.resize(640, 480)

        self._build_ui()
        self._build_menus()
        self._load_preferences()
        self.statusBar().showMessage("Listo")

    def _build_ui(self) -> None:
        central = QWidget()
        layout = QVBoxLayout(central)

        button_row = QHBoxLayout()
        self.open_button = QPushButton("Abrir Excel...")
        self.open_button.clicked.connect(self.open_excel)
        button_row.addWidget(self.open_button)

        self.convert_button = QPushButton("Convertir")
        self.convert_button.clicked.connect(self.convert)
        button_row.addWidget(self.convert_button)

        self.save_button = QPushButton("Guardar SQL...")
        self.save_button.clicked.connect(self.save_sql)
        button_row.addWidget(self.save_button)

        layout.addLayout(button_row)

        self.editor = QPlainTextEdit()
        layout.addWidget(self.editor)

        self.setCentralWidget(central)

    def _build_menus(self) -> None:
        menu_bar = self.menuBar()

        file_menu = menu_bar.addMenu("Archivo")

        open_action = QAction(_icon("open"), "Abrir Excel...", self)
        open_action.setShortcut("Ctrl+O")
        open_action.triggered.connect(self.open_excel)
        file_menu.addAction(open_action)

        convert_action = QAction(_icon("build"), "Convertir", self)
        convert_action.setShortcut("Ctrl+R")
        convert_action.triggered.connect(self.convert)
        file_menu.addAction(convert_action)

        save_action = QAction(_icon("save"), "Guardar SQL...", self)
        save_action.setShortcut("Ctrl+S")
        save_action.triggered.connect(self.save_sql)
        file_menu.addAction(save_action)

        file_menu.addSeparator()

        exit_action = QAction(_icon("exit_to_app"), "Salir", self)
        exit_action.setShortcut("Ctrl+Q")
        exit_action.triggered.connect(self.close)
        file_menu.addAction(exit_action)

        options_menu = menu_bar.addMenu("Opciones")

        self.create_table_action = QAction("Incluir CREATE TABLE", self)
        self.create_table_action.setCheckable(True)
        self.create_table_action.setChecked(self.include_create)
        self.create_table_action.toggled.connect(self._set_include_create)
        options_menu.addAction(self.create_table_action)

        view_menu = menu_bar.addMenu("Ver")
        self.dark_mode_action = QAction(_icon("brightness_6"), "Alternar modo oscuro/claro", self)
        self.dark_mode_action.triggered.connect(self.toggle_dark_mode)
        view_menu.addAction(self.dark_mode_action)

        help_menu = menu_bar.addMenu("Ayuda")
        instructions_action = QAction(_icon("help"), "Instrucciones", self)
        instructions_action.setShortcut("F1")
        instructions_action.triggered.connect(self.show_instructions)
        help_menu.addAction(instructions_action)

        about_action = QAction(_icon("info"), "Acerca de...", self)
        about_action.triggered.connect(self.show_about)
        help_menu.addAction(about_action)

    def open_excel(self) -> None:
        filename, _ = QFileDialog.getOpenFileName(
            self, "Abrir Excel", "", "Archivos Excel (*.xlsx *.xlsm)"
        )
        if not filename:
            return
        self.excel_path = Path(filename)
        self._choose_sheet_and_convert()

    def convert(self) -> None:
        if self.excel_path is None:
            self.open_excel()
            return
        self._choose_sheet_and_convert()

    def _choose_sheet_and_convert(self) -> None:
        try:
            sheet_names = list_sheet_names(str(self.excel_path))
        except Exception as exc:  # noqa: BLE001 - mostrar cualquier error al usuario
            QMessageBox.critical(self, "Error", f"No se pudo leer el archivo:\n{exc}")
            return

        if not sheet_names:
            QMessageBox.information(
                self, "Sin tablas", "El archivo no contiene hojas para convertir."
            )
            return

        dialog = SheetSelectionDialog(sheet_names, self, dark_mode=self.dark_mode)
        if dialog.exec() != QDialog.DialogCode.Accepted or dialog.selected_sheet is None:
            self.statusBar().showMessage("Conversión cancelada")
            return

        try:
            statements = generate_insert_statements(
                str(self.excel_path),
                [dialog.selected_sheet],
                include_create=self.include_create,
            )
        except Exception as exc:  # noqa: BLE001 - mostrar cualquier error al usuario
            QMessageBox.critical(self, "Error", f"No se pudo convertir el archivo:\n{exc}")
            return

        self.editor.setPlainText("\n".join(statements))
        self.statusBar().showMessage(
            f"{len(statements)} sentencias generadas desde "
            f"{dialog.selected_sheet} de {self.excel_path.name}"
        )

    def save_sql(self) -> None:
        filename, _ = QFileDialog.getSaveFileName(
            self, "Guardar SQL", "insert_statements.sql", "Archivos SQL (*.sql)"
        )
        if not filename:
            return
        try:
            Path(filename).write_text(self.editor.toPlainText(), encoding="utf-8")
        except OSError as exc:
            QMessageBox.critical(self, "Error", f"No se pudo guardar:\n{exc}")
            return
        self.statusBar().showMessage(f"Guardado en {filename}")

    def toggle_dark_mode(self) -> None:
        self._apply_dark_mode(not self.dark_mode)

    def _set_include_create(self, enabled: bool) -> None:
        self.include_create = enabled

    def show_instructions(self) -> None:
        HelpDialog(self, dark_mode=self.dark_mode).exec()

    def show_about(self) -> None:
        AboutDialog(self, dark_mode=self.dark_mode).exec()

    def _apply_dark_mode(self, enabled: bool) -> None:
        self.dark_mode = enabled
        self.setStyleSheet(DARK_STYLE if enabled else "")

    def _center(self) -> None:
        screen = QApplication.primaryScreen()
        if screen is None:
            return
        frame = self.frameGeometry()
        frame.moveCenter(screen.availableGeometry().center())
        self.move(frame.topLeft())

    def _load_preferences(self) -> None:
        if not CONFIG_FILE.exists():
            self._center()
            return
        try:
            root = ET.parse(CONFIG_FILE).getroot()
        except ET.ParseError:
            self._center()
            return

        window = root.find("window")
        if window is None:
            self._center()
        else:
            try:
                self.resize(int(window.findtext("width")), int(window.findtext("height")))
                self.move(int(window.findtext("x")), int(window.findtext("y")))
            except (TypeError, ValueError):
                self._center()

        dark_mode = root.findtext("settings/dark_mode")
        if dark_mode is not None:
            self._apply_dark_mode(dark_mode.lower() == "true")

        include_create = root.findtext("settings/include_create")
        if include_create is not None:
            self._set_include_create(include_create.lower() == "true")
            self.create_table_action.setChecked(self.include_create)

    def _save_preferences(self) -> None:
        root = ET.Element(PROGRAM_NAME)
        window = ET.SubElement(root, "window")
        ET.SubElement(window, "x").text = str(self.x())
        ET.SubElement(window, "y").text = str(self.y())
        ET.SubElement(window, "width").text = str(self.width())
        ET.SubElement(window, "height").text = str(self.height())
        settings = ET.SubElement(root, "settings")
        ET.SubElement(settings, "dark_mode").text = str(self.dark_mode).lower()
        ET.SubElement(settings, "include_create").text = str(self.include_create).lower()
        tree = ET.ElementTree(root)
        ET.indent(tree, space="  ")
        CONFIG_FILE.parent.mkdir(parents=True, exist_ok=True)
        tree.write(CONFIG_FILE, encoding="utf-8", xml_declaration=True)

    def closeEvent(self, event) -> None:
        self._save_preferences()
        event.accept()


def main() -> int:
    app = QApplication(sys.argv)
    app.setApplicationName(PROGRAM_NAME)
    app.setWindowIcon(_logo_icon())
    window = MainWindow()
    window.show()
    return app.exec()


if __name__ == "__main__":
    raise SystemExit(main())
